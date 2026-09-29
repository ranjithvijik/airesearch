"""Build shareable prediction-report artifacts from the consolidated manuscript.

Reads article/ai_circularity_article.tex and produces, with pandoc:
  AI_Bubble_Burst_Prediction_Report.docx  (styles inherited from the
      technical-report .docx so both reports look like siblings;
      refreshable TOC field, one part per page, page numbers, H1 parts
      with H2 sections underneath)
  prediction.html  (standalone, embedded figures, MathJax, part-grouped
      clickable TOC, part banners, collapsible appendices, back-to-top
      links, dashboard-echo stylesheet)

One preprocessing step the .tex needs for pandoc: natbib \\citet/\\citep
commands are rewritten to literal APA author-date strings derived from the
paper's own bibliography (no external .bib), so citations survive in both
outputs. Section labels/refs, tables, and figures pass through natively.
python-docx demotes pandoc's Heading 3 sections to Heading 2 (parts stay H1).
BeautifulSoup reorganizes the HTML DOM (heading levels, grouped TOC,
appendix <details>, part banners).

Requires: pandoc>=3, python-docx, beautifulsoup4.
Run: python3 build_prediction_artifacts.py
"""
import os
import re
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.abspath(__file__))
TEX = os.path.join(ROOT, "article", "ai_circularity_article.tex")
REFDOC = os.path.join(ROOT, "AI_Ecosystem_Game_Theory_Technical_Report.docx")
OUT_DOCX = os.path.join(ROOT, "AI_Bubble_Burst_Prediction_Report.docx")
OUT_HTML = os.path.join(ROOT, "prediction.html")

CSS = """body{font-family:"Inter","SF Pro Display",-apple-system,"Segoe UI",Roboto,"Helvetica Neue",Arial,sans-serif;max-width:1000px;margin:0 auto;padding:24px 28px;color:#1B2A3A;line-height:1.6;background:#F7F9FA}
header{border-bottom:3px solid #0072B2;margin-bottom:20px}
#TOC{background:#fff;border:1px solid #DDE3E8;border-radius:12px;padding:14px 20px;box-shadow:0 2px 8px rgba(27,42,58,.07);margin-bottom:22px}
.toc-title{font-weight:700;font-size:16px;margin:0 0 8px;color:#1B2A3A}
.toc-groups{columns:2;column-gap:28px}.toc-group{break-inside:avoid;margin-bottom:10px}
.toc-part{font-weight:700;margin:0 0 2px;font-size:14px}
.toc-group ul{margin:2px 0 6px 18px;padding:0}.toc-group li{margin:2px 0;font-size:13.5px}
h1.part{color:#fff;background:linear-gradient(120deg,#1B2A3A,#14425F);padding:12px 18px;border-radius:10px;font-size:22px}
.part-eyebrow{display:block;font-size:12px;font-weight:600;letter-spacing:2px;text-transform:uppercase;color:#8FD0F2;margin-bottom:4px}
h1:not(.part){color:#1B2A3A;border-bottom:2px solid #0072B2;padding-bottom:6px}
h2{color:#14425F}h3{color:#44566A}
details.appendix{background:#fff;border:1px solid #DDE3E8;border-radius:12px;padding:10px 16px;margin:16px 0;box-shadow:0 2px 8px rgba(27,42,58,.07)}
details.appendix summary{cursor:pointer;font-weight:700;color:#14425F}
details.appendix summary h2{display:inline;font-size:18px;margin:0}
.totop{text-align:right;font-size:13px;margin:6px 0 0}
img{max-width:100%!important;height:auto!important}
figure{background:#fff;border:1px solid #DDE3E8;border-radius:12px;padding:14px;box-shadow:0 2px 8px rgba(27,42,58,.07);margin:18px 0;text-align:center}
figure img{max-width:100%;height:auto}figcaption{font-size:13px;color:#5D6D7E;margin-top:8px;text-align:left}
table{border-collapse:collapse;width:100%;margin:14px 0;font-size:14px;background:#fff}
th,td{border:1px solid #DDE3E8;padding:7px 10px;text-align:left}
th{background:#E7ECEF}tr:nth-child(even){background:#F8F9F9}
caption{font-weight:700;margin-bottom:6px;text-align:left}
code{background:#EEF1F4;padding:1px 5px;border-radius:4px;font-size:.9em}
a{color:#0072B2}blockquote{border-left:4px solid #0072B2;background:#E8F4FD;margin:14px 0;padding:10px 14px;border-radius:0 8px 8px 0}
#refs div{padding-left:1.5em;text-indent:-1.5em;margin-bottom:6px}
@media print{#TOC{columns:1}}
"""


def apa_short(optional_label):
    """'Agrawal et al.(2019)' -> 'Agrawal et al. (2019)'."""
    m = re.match(r"(.*)\((\d{4}[a-z]?)\)(.*)$", optional_label.strip())
    if not m:
        return optional_label.strip()
    return (m.group(1) + " (" + m.group(2) + ")").replace("et al ",
                                                          "et al. ").replace("et al..", "et al.")


def rewrite_citations(src):
    keys = {}

    def bib(m):
        keys[m.group(2)] = apa_short(m.group(1))
        return m.group(0)
    src = re.sub(r"\\bibitem\[(.*?)\]\{(.*?)\}", bib, src)

    def cite(pattern, fmt):
        def sub(m):
            parts = []
            for k in m.group(1).split(","):
                k = k.strip()
                parts.append(keys.get(k, k))
            return fmt(parts)
        return re.sub(pattern, sub, src)
    src = cite(r"\\citet\{([^}]*)\}", "; ".join)
    src = cite(r"\\citep\{([^}]*)\}",
               lambda parts: "(" + "; ".join(parts) + ")")
    return src, keys


def run(cmd, **kw):
    r = subprocess.run(cmd, capture_output=True, text=True, **kw)
    if r.returncode != 0:
        print(r.stdout[-2000:])
        print(r.stderr[-2000:])
        raise SystemExit(f"FAILED: {' '.join(cmd)}")
    return r


def organize_docx(path, has_parts):
    """Normalize heading levels, page break before each part, page numbers.

    With \\parts: H3 sections -> H2, parts stay H1 with page breaks.
    Without (flat article layout): H1 sections -> H2, H2 subsections -> H3.
    """
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = docx.Document(path)
    n_demoted = 0
    if has_parts:
        for p in doc.paragraphs:
            if p.style.name == "Heading 3":
                p.style = doc.styles["Heading 2"]
                n_demoted += 1
    else:
        for p in doc.paragraphs:
            if p.style.name == "Heading 1":
                p.style = doc.styles["Heading 2"]
                n_demoted += 1
            elif p.style.name == "Heading 2":
                p.style = doc.styles["Heading 3"]
                n_demoted += 1
    n_breaks = 0
    for p in doc.paragraphs:
        if has_parts and p.style.name == "Heading 1":
            pb = p.insert_paragraph_before()
            pb.add_run().add_break(WD_BREAK.PAGE)
            n_breaks += 1
    sect = doc.sections[0]
    footer = sect.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar" if text is None else "w:instrText")
        if text is None:
            el.set(qn("w:fldCharType"), tag)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    doc.save(path)
    print(f"docx organized: {n_demoted} sections demoted to H2, "
          f"{n_breaks} part page-breaks, page-number footer added")


def organize_html(path):
    """Fix heading levels, part banners, grouped TOC, appendix details."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(open(path, encoding="utf-8"), "html.parser")
    for h in soup.find_all("h3"):
        h.name = "h2"
    romans = ["I", "II", "III", "IV", "V", "VI", "VII"]
    parts = [h for h in soup.find_all("h1")
             if (h.get("id") or "").startswith("part:")]
    flat = not parts
    if flat:
        # No parts: pandoc promotes sections to h1; restore h2/h3 nesting.
        for h in soup.find_all("h1"):
            if "title" not in h.get("class", []):
                h.name = "h2"
        for h in soup.find_all("h3"):
            h.name = "h2"
        for h in soup.find_all("h2"):
            if (h.get("id") or "") in ("sec:howtoread", "sec:policy-bcr"):
                h.name = "h3"

    def own_text(h):
        return " ".join(t.strip() for t in h.find_all(string=True, recursive=False)
                        if t.strip())

    for i, h in enumerate(parts):
        h["class"] = list(h.get("class", [])) + ["part"]
        eye = soup.new_tag("span", **{"class": "part-eyebrow"})
        eye.string = f"Part {romans[i]}" if i < len(romans) else f"Part {i + 1}"
        h.insert(0, eye)

    toc = soup.find("nav", id="TOC")
    n_links = 0
    if toc is not None:
        toc.clear()
        title = soup.new_tag("p", **{"class": "toc-title"})
        title.string = "Contents"
        toc.append(title)
        groups = soup.new_tag("div", **{"class": "toc-groups"})
        toc.append(groups)
        if not parts:
            g = soup.new_tag("div", **{"class": "toc-group"})
            groups.append(g)
            ul = soup.new_tag("ul")
            g.append(ul)
            for sib in soup.find_all("h2"):
                if sib.get("id"):
                    li = soup.new_tag("li")
                    sa = soup.new_tag("a", href="#" + sib["id"])
                    sa.string = own_text(sib) or sib.get_text(strip=True)
                    li.append(sa)
                    ul.append(li)
                    n_links += 1
        for h in parts:
            g = soup.new_tag("div", **{"class": "toc-group"})
            groups.append(g)
            pp = soup.new_tag("p", **{"class": "toc-part"})
            a = soup.new_tag("a", href="#" + h["id"])
            a.string = own_text(h) or h.get_text(strip=True)
            pp.append(a)
            g.append(pp)
            ul = soup.new_tag("ul")
            g.append(ul)
            sib = h.next_sibling
            while sib is not None and getattr(sib, "name", None) != "h1":
                if getattr(sib, "name", None) == "h2" and sib.get("id"):
                    li = soup.new_tag("li")
                    sa = soup.new_tag("a", href="#" + sib["id"])
                    sa.string = own_text(sib) or sib.get_text(strip=True)
                    li.append(sa)
                    ul.append(li)
                    n_links += 1
                sib = sib.next_sibling
            if not ul.find_all("li"):
                ul.decompose()

    n_app = 0
    for h in soup.find_all("h2"):
        if (h.get("id") or "").startswith("sec:appendix-"):
            det = soup.new_tag("details", **{"class": "appendix"})
            det["open"] = ""
            h.insert_before(det)
            summ = soup.new_tag("summary")
            det.append(summ)
            summ.append(h.extract())
            sib = det.next_sibling
            while sib is not None and getattr(sib, "name", None) not in ("h1", "h2"):
                nxt = sib.next_sibling
                det.append(sib.extract())
                sib = nxt
            n_app += 1

    def add_totop(anchor):
        p = soup.new_tag("p", **{"class": "totop"})
        a = soup.new_tag("a", href="#TOC")
        a.string = "Back to top \u2191"
        p.append(a)
        anchor.insert_before(p)

    n_top = 0
    if flat:
        secs = [h for h in soup.find_all("h2")
                if h.get("id") and not h.find_parent("details")]
        for h in secs[5::6]:
            add_totop(h)
            n_top += 1
    else:
        for h in parts[1:]:
            add_totop(h)
            n_top += 1
    if soup.body is not None:
        p = soup.new_tag("p", **{"class": "totop"})
        a = soup.new_tag("a", href="#TOC")
        a.string = "Back to top \u2191"
        p.append(a)
        soup.body.append(p)
    open(path, "w", encoding="utf-8").write(str(soup))
    print(f"html organized: {len(parts)} part banners, {n_links} grouped TOC links, "
          f"{n_app} appendix <details>, {n_top + 1} back-to-top links")


def main():
    src = open(TEX, encoding="utf-8").read()
    src, keys = rewrite_citations(src)
    print(f"citations mapped: {len(keys)} keys")
    assert len(keys) == 43, f"expected 43 bib keys, got {len(keys)}"
    tmp = tempfile.mkdtemp(prefix="pred_artifacts_")
    pre = os.path.join(tmp, "prediction_pandoc.tex")
    open(pre, "w", encoding="utf-8").write(src)

    run(["pandoc", pre, "-o", OUT_DOCX, "--reference-doc=" + REFDOC,
         "--toc", "--toc-depth=2"],
        cwd=os.path.join(ROOT, "article"))
    organize_docx(OUT_DOCX, has_parts="\\part{" in src)

    css = os.path.join(tmp, "prediction.css")
    open(css, "w", encoding="utf-8").write(CSS)
    run(["pandoc", pre, "-s", "--toc", "--toc-depth=2", "--mathjax",
         "--embed-resources", "--css=" + css,
         "-o", OUT_HTML], cwd=os.path.join(ROOT, "article"))
    organize_html(OUT_HTML)
    print(f"wrote {OUT_HTML} ({os.path.getsize(OUT_HTML) / 1024:.0f} KB)")
    print(f"wrote {OUT_DOCX} ({os.path.getsize(OUT_DOCX) / 1024:.0f} KB)")


if __name__ == "__main__":
    main()